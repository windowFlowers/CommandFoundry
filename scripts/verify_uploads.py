from __future__ import annotations

import json
import os
import sys
import time
from io import BytesIO

import httpx
from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


API_BASE = os.getenv("AEGIS_VERIFY_API", "http://127.0.0.1:8002").rstrip("/")
KNOWLEDGE_BASE_NAME = os.getenv("AEGIS_VERIFY_KB_NAME", "格式验收知识库")

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


def docx_bytes() -> bytes:
    buffer = BytesIO()
    document = DocxDocument()
    document.add_heading("Python virtual environment", level=1)
    document.add_paragraph("Create an isolated environment before installing dependencies.")
    document.add_paragraph("python -m venv .venv")
    document.save(buffer)
    return buffer.getvalue()


def pdf_bytes() -> bytes:
    buffer = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 72 720 Td (Kubernetes rollout status check) Tj "
        b"0 -22 Td (kubectl rollout status deployment/api) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(buffer)
    return buffer.getvalue()


def wait_until_ready(client: httpx.Client, document_id: str) -> dict:
    deadline = time.monotonic() + 45
    while time.monotonic() < deadline:
        response = client.get(f"/knowledge-documents/{document_id}")
        response.raise_for_status()
        document = response.json()
        if document["status"] == "ready":
            return document
        if document["status"] == "failed":
            raise RuntimeError(f"Indexing failed for {document['filename']}: {document.get('error')}")
        time.sleep(0.2)
    raise TimeoutError(f"Indexing timed out for document {document_id}")


def main() -> None:
    fixtures = [
        ("powershell-health.txt", b"PowerShell service health check\n\nGet-Service -Name Spooler", "Get-Service"),
        (
            "node-runbook.md",
            b"# Node dependency audit\n\nRun before release.\n\n```bash\nnpm audit --omit=dev\n```",
            "npm audit",
        ),
        ("kubernetes-rollout.pdf", pdf_bytes(), "kubectl rollout status"),
        ("python-venv.docx", docx_bytes(), "python -m venv"),
    ]
    with httpx.Client(base_url=API_BASE, timeout=30) as client:
        bases = client.get("/knowledge-bases").json()["items"]
        for knowledge_base in bases:
            if knowledge_base["name"] == KNOWLEDGE_BASE_NAME and not knowledge_base["is_builtin"]:
                response = client.delete(f"/knowledge-bases/{knowledge_base['id']}")
                response.raise_for_status()

        response = client.post("/knowledge-bases", json={"name": KNOWLEDGE_BASE_NAME})
        response.raise_for_status()
        knowledge_base = response.json()
        indexed = []
        for filename, payload, marker in fixtures:
            response = client.post(
                f"/knowledge-bases/{knowledge_base['id']}/documents",
                files={"file": (filename, payload, "application/octet-stream")},
            )
            if response.status_code != 202:
                raise RuntimeError(f"Upload failed for {filename}: {response.status_code} {response.text}")
            document = wait_until_ready(client, response.json()["id"])
            preview = client.get(f"/knowledge-documents/{document['id']}/content")
            preview.raise_for_status()
            if marker not in preview.json()["content"]:
                raise AssertionError(f"Preview marker is missing for {filename}")
            indexed.append(document)

        duplicate = client.post(
            f"/knowledge-bases/{knowledge_base['id']}/documents",
            files={"file": (fixtures[0][0], fixtures[0][1], "text/plain")},
        )
        if duplicate.status_code != 409:
            raise AssertionError(f"Expected duplicate upload to return 409, got {duplicate.status_code}")

        response = client.post(f"/knowledge-documents/{indexed[0]['id']}/reindex")
        if response.status_code != 202:
            raise RuntimeError(f"Reindex failed: {response.status_code} {response.text}")
        indexed[0] = wait_until_ready(client, indexed[0]["id"])
        status = client.get(
            "/knowledge/status", params={"knowledge_base_id": knowledge_base["id"]}
        )
        status.raise_for_status()

    print(
        json.dumps(
            {
                "knowledge_base_id": knowledge_base["id"],
                "knowledge_base_name": knowledge_base["name"],
                "documents": [
                    {
                        "filename": document["filename"],
                        "status": document["status"],
                        "chunk_count": document["chunk_count"],
                    }
                    for document in indexed
                ],
                "duplicate_rejected": True,
                "reindex_verified": True,
                "status": status.json(),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()

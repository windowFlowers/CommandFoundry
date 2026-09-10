from __future__ import annotations

import hashlib
from io import BytesIO

from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.chunking import build_parent_child_chunks, extract_command_evidence
from app.extraction import ExtractedDocument, ExtractionService


def _sha(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _pdf_bytes() -> bytes:
    output = BytesIO()
    writer = PdfWriter()
    for text in (b"First page procedure", b"Second page verification"):
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        font_ref = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 72 720 Td (" + text + b") Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(output)
    return output.getvalue()


def test_markdown_blocks_have_heading_lines_and_exact_character_ranges() -> None:
    source = b"# Deploy\r\n\r\nCheck state.\r\n\r\n## Run\r\n\r\n`npm test`\r\n"
    document = ExtractionService().extract("guide.md", source)

    assert isinstance(document, ExtractedDocument)
    assert isinstance(document, str)  # v2.2 compatibility
    assert [block.kind for block in document.blocks] == ["heading", "paragraph", "heading", "paragraph"]
    assert document.blocks[1].heading_path == ("Deploy",)
    assert document.blocks[2].heading_path == ("Deploy", "Run")
    assert document.blocks[2].line_start == document.blocks[2].line_end == 5
    for block in document.blocks:
        assert document.content[block.char_start : block.char_end] == block.text


def test_txt_and_pdf_locators_are_one_based() -> None:
    text_document = ExtractionService().extract("notes.txt", b"first\nline\n\nthird")
    assert text_document.blocks[0].line_start == 1
    assert text_document.blocks[0].line_end == 2
    assert text_document.blocks[1].line_start == text_document.blocks[1].line_end == 4

    pdf_document = ExtractionService().extract("pages.pdf", _pdf_bytes())
    assert [block.page_start for block in pdf_document.blocks] == [1, 2]
    assert all(block.paragraph_start == 1 for block in pdf_document.blocks)
    assert "Second page verification" in pdf_document.content


def test_docx_preserves_interleaved_paragraph_and_table_order() -> None:
    output = BytesIO()
    document = DocxDocument()
    document.add_heading("Operations", level=1)
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Key"
    table.cell(0, 1).text = "Value"
    document.add_paragraph("After table")
    document.save(output)

    extracted = ExtractionService().extract("operations.docx", output.getvalue())
    values = [block.text for block in extracted.blocks]
    assert values == ["Operations", "Before table", "Key | Value", "After table"]
    assert extracted.blocks[2].kind == "table_row"
    assert extracted.blocks[2].table_start == extracted.blocks[2].table_row_start == 1
    assert extracted.blocks[-1].heading_path == ("Operations",)


def test_parent_child_ids_are_stable_and_heading_is_search_only_prefix() -> None:
    source = (
        "# Deployment\n\n"
        + "Preparation sentence. " * 40
        + "\n\n## Verification\n\n"
        + "Health verification detail. " * 50
    ).encode()
    extracted = ExtractionService().extract("deploy.md", source)
    first = build_parent_child_chunks(
        extracted,
        _sha(source),
        parent_target=360,
        parent_max=480,
        child_target=120,
        child_max=180,
        child_overlap=30,
    )
    second = build_parent_child_chunks(
        extracted,
        _sha(source),
        parent_target=360,
        parent_max=480,
        child_target=120,
        child_max=180,
        child_overlap=30,
    )

    assert [item.id for item in first.parents] == [item.id for item in second.parents]
    assert [item.id for item in first.children] == [item.id for item in second.children]
    assert all(len(item.text) <= 480 for item in first.parents)
    assert all(len(item.text) <= 180 for item in first.children)
    detail = next(item for item in first.children if "Health verification detail" in item.text)
    assert "Deployment > Verification" in detail.retrieval_text
    assert "Deployment > Verification" not in detail.text


def test_long_fence_is_split_only_on_lines_and_each_piece_is_reclosed() -> None:
    code_lines = [f"echo line-{index:03d}" for index in range(80)]
    source = ("# Script\n\n```bash\n" + "\n".join(code_lines) + "\n```\n").encode()
    extracted = ExtractionService().extract("script.md", source)
    chunks = build_parent_child_chunks(
        extracted,
        _sha(source),
        parent_target=500,
        parent_max=650,
        child_target=140,
        child_max=180,
        child_overlap=20,
    )
    code_chunks = [item for item in chunks.children if "echo line-" in item.text]

    assert len(code_chunks) > 1
    assert all(item.text.count("```") == 2 for item in code_chunks)
    assert all(len(item.text) <= 180 for item in code_chunks)
    recovered = [
        line
        for item in code_chunks
        for line in item.command_evidence[0].code.splitlines()
        if line.startswith("echo line-")
    ]
    assert recovered == code_lines


def test_command_evidence_is_limited_to_explicit_code_shapes() -> None:
    text = """This paragraph discusses deployment but is not a command.
`npm test`
docker compose restart worker
```powershell
Get-Service -Name Spooler
```"""
    evidence = extract_command_evidence(text)

    assert {(item.kind, item.code) for item in evidence} == {
        ("inline_code", "npm test"),
        ("standalone_command", "docker compose restart worker"),
        ("fenced_code", "Get-Service -Name Spooler"),
    }

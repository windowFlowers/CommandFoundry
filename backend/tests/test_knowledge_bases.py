from __future__ import annotations

import sqlite3
import time
from io import BytesIO
from datetime import UTC, datetime
from pathlib import Path

from fastapi.testclient import TestClient
from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.main import app
from app.models import DEFAULT_KNOWLEDGE_BASE_ID
from app.repository import ConversationRepository
from app.text import split_into_chunks


def _wait_for_document(client: TestClient, document_id: str, expected: str = "ready") -> dict:
    for _ in range(100):
        payload = client.get(f"/knowledge-documents/{document_id}").json()
        if payload["status"] in {expected, "failed"}:
            return payload
        time.sleep(0.02)
    raise AssertionError("document indexing did not finish")


def _docx_bytes() -> bytes:
    buffer = BytesIO()
    document = DocxDocument()
    document.add_heading("Python virtual environment", level=1)
    document.add_paragraph("Create an isolated environment before installing dependencies.")
    document.add_paragraph("python -m venv .venv")
    document.save(buffer)
    return buffer.getvalue()


def _pdf_bytes() -> bytes:
    buffer = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 72 720 Td (Kubernetes rollout status check) Tj "
        b"0 -22 Td (kubectl rollout status deployment/api) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(buffer)
    return buffer.getvalue()


def test_knowledge_base_crud_upload_preview_and_isolation() -> None:
    with TestClient(app) as client:
        builtins = client.get("/knowledge-bases").json()["items"]
        assert builtins[0]["id"] == DEFAULT_KNOWLEDGE_BASE_ID
        assert builtins[0]["is_builtin"] is True
        assert client.delete(f"/knowledge-bases/{DEFAULT_KNOWLEDGE_BASE_ID}").status_code == 409

        created = client.post("/knowledge-bases", json={"name": "团队运行手册"}).json()
        knowledge_base_id = created["id"]
        renamed = client.patch(
            f"/knowledge-bases/{knowledge_base_id}", json={"name": "后端运行手册"}
        ).json()
        assert renamed["name"] == "后端运行手册"

        source = b"""# Queue operations

Restart a worker after checking the queue depth.

```bash
docker compose restart worker
```
"""
        uploaded = client.post(
            f"/knowledge-bases/{knowledge_base_id}/documents",
            files={"file": ("runbook.md", source, "text/markdown")},
        )
        assert uploaded.status_code == 202
        document = _wait_for_document(client, uploaded.json()["id"])
        assert document["status"] == "ready"
        assert document["chunk_count"] >= 1
        content = client.get(f"/knowledge-documents/{document['id']}/content").json()
        assert "docker compose restart worker" in content["content"]

        duplicate = client.post(
            f"/knowledge-bases/{knowledge_base_id}/documents",
            files={"file": ("copy.md", source, "text/markdown")},
        )
        assert duplicate.status_code == 409

        conversation = client.post(
            "/conversations", json={"title": "队列排查", "knowledge_base_id": knowledge_base_id}
        ).json()
        assert conversation["knowledge_base_id"] == knowledge_base_id
        mismatch = client.post(
            "/chat/stream",
            json={
                "query": "怎么重启 worker？",
                "conversation_id": conversation["id"],
                "knowledge_base_id": DEFAULT_KNOWLEDGE_BASE_ID,
            },
        )
        assert mismatch.status_code == 409

        assert client.delete(f"/knowledge-bases/{knowledge_base_id}").status_code == 204
        historical = client.get(f"/conversations/{conversation['id']}").json()
        assert historical["knowledge_base_deleted"] is True
        assert client.post(
            "/chat/stream", json={"query": "继续", "conversation_id": conversation["id"]}
        ).status_code == 409


def test_upload_validation_and_failed_empty_document() -> None:
    with TestClient(app) as client:
        unsupported = client.post(
            f"/knowledge-bases/{DEFAULT_KNOWLEDGE_BASE_ID}/documents",
            files={"file": ("script.exe", b"binary", "application/octet-stream")},
        )
        assert unsupported.status_code == 400
        empty = client.post(
            f"/knowledge-bases/{DEFAULT_KNOWLEDGE_BASE_ID}/documents",
            files={"file": ("empty.txt", b"", "text/plain")},
        )
        assert empty.status_code == 400


def test_all_supported_document_formats_extract_index_and_preview() -> None:
    fixtures = [
        ("health.txt", b"PowerShell health check\n\nGet-Service -Name Spooler", "Get-Service"),
        (
            "audit.md",
            b"# Node audit\n\n```bash\nnpm audit --omit=dev\n```",
            "npm audit",
        ),
        ("rollout.pdf", _pdf_bytes(), "kubectl rollout status"),
        ("venv.docx", _docx_bytes(), "python -m venv"),
    ]
    with TestClient(app) as client:
        knowledge_base = client.post("/knowledge-bases", json={"name": "格式解析验收"}).json()
        for filename, payload, marker in fixtures:
            response = client.post(
                f"/knowledge-bases/{knowledge_base['id']}/documents",
                files={"file": (filename, payload, "application/octet-stream")},
            )
            assert response.status_code == 202
            document = _wait_for_document(client, response.json()["id"])
            assert document["status"] == "ready"
            preview = client.get(f"/knowledge-documents/{document['id']}/content").json()
            assert marker in preview["content"]


def test_code_aware_chunker_never_splits_fenced_block() -> None:
    fenced = "```bash\n" + "echo safe\n" * 180 + "```"
    text = "# Deploy\n\nBefore deploy, inspect state.\n\n" + fenced + "\n\nAfter deploy, verify health."
    chunks = split_into_chunks(text, chunk_size=240, overlap=40)
    containing = [chunk for chunk in chunks if "```bash" in chunk]
    assert len(containing) == 1
    assert containing[0].count("```") == 2
    assert "echo safe" in containing[0]


def test_legacy_database_migrates_conversations_to_builtin_knowledge_base(tmp_path: Path) -> None:
    database = tmp_path / "legacy.sqlite3"
    now = datetime.now(UTC).isoformat()
    with sqlite3.connect(database) as connection:
        connection.executescript(
            """
            CREATE TABLE conversations (id TEXT PRIMARY KEY, title TEXT NOT NULL, created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
            CREATE TABLE messages (id TEXT PRIMARY KEY, conversation_id TEXT NOT NULL, role TEXT NOT NULL, query TEXT, answer_json TEXT, created_at TEXT NOT NULL);
            """
        )
        connection.execute(
            "INSERT INTO conversations(id, title, created_at, updated_at) VALUES ('legacy', '旧会话', ?, ?)",
            (now, now),
        )
    repository = ConversationRepository(database)
    conversation = repository.get("legacy")
    assert conversation is not None
    assert conversation.knowledge_base_id == DEFAULT_KNOWLEDGE_BASE_ID
    assert conversation.knowledge_base_deleted is False

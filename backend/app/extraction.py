from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from io import BytesIO
from pathlib import Path
from typing import Literal

try:
    from pypdf import PdfReader
except ModuleNotFoundError:  # pragma: no cover - optional import is verified at packaging time
    PdfReader = None

try:
    from docx import Document as DocxDocument
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ModuleNotFoundError:  # pragma: no cover
    DocxDocument = None
    CT_Tbl = CT_P = Table = Paragraph = None


SourceKind = Literal["markdown", "text", "pdf", "docx"]
BlockKind = Literal["heading", "paragraph", "code", "table_row"]

_MARKDOWN_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_DOCX_HEADING = re.compile(r"(?:heading|标题)\s*([1-6])$", re.IGNORECASE)


class ExtractionError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class ExtractedBlock:
    """One ordered structural element in the canonical extracted text.

    Character offsets are zero-based, end-exclusive offsets into
    :attr:`ExtractedDocument.content`. Human-facing ordinals are one-based.
    """

    text: str
    kind: BlockKind
    char_start: int
    char_end: int
    heading_path: tuple[str, ...] = ()
    page_start: int | None = None
    page_end: int | None = None
    line_start: int | None = None
    line_end: int | None = None
    paragraph_start: int | None = None
    paragraph_end: int | None = None
    table_start: int | None = None
    table_end: int | None = None
    table_row_start: int | None = None
    table_row_end: int | None = None
    element_start: int | None = None
    element_end: int | None = None

    def locator_dict(self, source_kind: SourceKind) -> dict:
        """Return a JSON-ready locator accepted by the API/model layer."""

        value = asdict(self)
        value.pop("text", None)
        value.pop("kind", None)
        value["kind"] = source_kind
        value["heading_path"] = list(self.heading_path)
        return value


class ExtractedDocument(str):
    """Structured extraction result that remains source-compatible with ``str``."""

    def __new__(
        cls,
        content: str,
        blocks: list[ExtractedBlock] | tuple[ExtractedBlock, ...],
        kind: SourceKind,
    ) -> "ExtractedDocument":
        instance = super().__new__(cls, content)
        instance.content = content
        instance.blocks = tuple(blocks)
        instance.kind = kind
        return instance

    content: str
    blocks: tuple[ExtractedBlock, ...]
    kind: SourceKind


@dataclass(frozen=True, slots=True)
class _PendingBlock:
    text: str
    kind: BlockKind
    heading_path: tuple[str, ...] = ()
    page_start: int | None = None
    page_end: int | None = None
    line_start: int | None = None
    line_end: int | None = None
    paragraph_start: int | None = None
    paragraph_end: int | None = None
    table_start: int | None = None
    table_end: int | None = None
    table_row_start: int | None = None
    table_row_end: int | None = None
    element_start: int | None = None
    element_end: int | None = None


def _normalize_lines(text: str) -> list[str]:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    return [line.rstrip() for line in text.split("\n")]


def _assemble(kind: SourceKind, pending: list[_PendingBlock]) -> ExtractedDocument:
    blocks: list[ExtractedBlock] = []
    content_parts: list[str] = []
    cursor = 0
    for item in pending:
        text = item.text.strip("\n")
        if not text.strip():
            continue
        if content_parts:
            content_parts.append("\n\n")
            cursor += 2
        start = cursor
        content_parts.append(text)
        cursor += len(text)
        blocks.append(
            ExtractedBlock(
                text=text,
                kind=item.kind,
                char_start=start,
                char_end=cursor,
                heading_path=item.heading_path,
                page_start=item.page_start,
                page_end=item.page_end,
                line_start=item.line_start,
                line_end=item.line_end,
                paragraph_start=item.paragraph_start,
                paragraph_end=item.paragraph_end,
                table_start=item.table_start,
                table_end=item.table_end,
                table_row_start=item.table_row_start,
                table_row_end=item.table_row_end,
                element_start=item.element_start,
                element_end=item.element_end,
            )
        )
    content = "".join(content_parts)
    if not content.strip():
        raise ExtractionError("文档没有可索引的文本")
    return ExtractedDocument(content, blocks, kind)


def _heading_level(paragraph: object) -> int | None:
    style = getattr(paragraph, "style", None)
    candidates = [getattr(style, "name", ""), getattr(style, "style_id", "")]
    for candidate in candidates:
        compact = re.sub(r"\s+", "", str(candidate or ""))
        match = re.search(r"(?:heading|标题)([1-6])$", compact, flags=re.IGNORECASE)
        if match:
            return int(match.group(1))
        match = _DOCX_HEADING.search(str(candidate or "").strip())
        if match:
            return int(match.group(1))
    return None


class ExtractionService:
    supported_suffixes = {".txt", ".md", ".markdown", ".pdf", ".docx"}

    def extract(self, filename: str, content: bytes) -> ExtractedDocument:
        suffix = Path(filename).suffix.lower()
        if suffix not in self.supported_suffixes:
            raise ExtractionError(f"不支持的文件类型：{suffix or '未知类型'}")
        if suffix in {".md", ".markdown"}:
            return self._extract_text_document(self._decode_text(content), markdown=True)
        if suffix == ".txt":
            return self._extract_text_document(self._decode_text(content), markdown=False)
        if suffix == ".pdf":
            return self._extract_pdf(content)
        return self._extract_docx(content)

    @staticmethod
    def _decode_text(content: bytes) -> str:
        # utf-8-sig must precede utf-8 or the BOM leaks into the first heading.
        for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise ExtractionError("无法解码文本文件内容")

    @staticmethod
    def _extract_text_document(text: str, *, markdown: bool) -> ExtractedDocument:
        lines = _normalize_lines(text)
        pending: list[_PendingBlock] = []
        heading_stack: list[str] = []
        current: list[str] = []
        current_start = 0
        in_fence = False
        fence_marker = ""

        def flush(kind: BlockKind = "paragraph") -> None:
            nonlocal current, current_start
            if not current:
                return
            start_offset = 0
            end_offset = len(current)
            while start_offset < end_offset and not current[start_offset].strip():
                start_offset += 1
            while end_offset > start_offset and not current[end_offset - 1].strip():
                end_offset -= 1
            if start_offset < end_offset:
                selected = current[start_offset:end_offset]
                pending.append(
                    _PendingBlock(
                        text="\n".join(selected),
                        kind=kind,
                        heading_path=tuple(heading_stack),
                        line_start=current_start + start_offset + 1,
                        line_end=current_start + end_offset,
                    )
                )
            current = []
            current_start = 0

        for line_index, line in enumerate(lines):
            stripped = line.strip()
            if in_fence:
                current.append(line)
                if stripped.startswith(fence_marker) and set(stripped) <= set(fence_marker):
                    in_fence = False
                    flush("code")
                    fence_marker = ""
                continue

            fence_match = re.match(r"^\s*(`{3,}|~{3,})", line)
            if fence_match:
                flush()
                current_start = line_index
                current = [line]
                fence_marker = fence_match.group(1)[0] * len(fence_match.group(1))
                in_fence = True
                continue

            heading_match = _MARKDOWN_HEADING.match(stripped) if markdown else None
            if heading_match:
                flush()
                level = len(heading_match.group(1))
                title = heading_match.group(2).strip()
                heading_stack[:] = heading_stack[: level - 1]
                heading_stack.append(title)
                pending.append(
                    _PendingBlock(
                        text=line,
                        kind="heading",
                        heading_path=tuple(heading_stack),
                        line_start=line_index + 1,
                        line_end=line_index + 1,
                    )
                )
                continue

            if not stripped:
                flush()
                continue
            if not current:
                current_start = line_index
            current.append(line)

        flush("code" if in_fence else "paragraph")
        return _assemble("markdown" if markdown else "text", pending)

    @staticmethod
    def _extract_pdf(content: bytes) -> ExtractedDocument:
        if PdfReader is None:
            raise ExtractionError("当前运行环境缺少 PDF 解析组件")
        try:
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ExtractionError("不支持加密 PDF")
            pending: list[_PendingBlock] = []
            for page_number, page in enumerate(reader.pages, start=1):
                page_text = page.extract_text() or ""
                # PDF libraries expose layout lines rather than semantic HTML.
                # These deterministic visual-line anchors are regrouped by the chunker.
                paragraphs = [line.strip() for line in _normalize_lines(page_text) if line.strip()]
                for paragraph_number, paragraph in enumerate(paragraphs, start=1):
                    pending.append(
                        _PendingBlock(
                            text=paragraph,
                            kind="paragraph",
                            page_start=page_number,
                            page_end=page_number,
                            paragraph_start=paragraph_number,
                            paragraph_end=paragraph_number,
                        )
                    )
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError("PDF 文件损坏或无法解析") from exc
        if not pending:
            raise ExtractionError("PDF 没有可提取文本，扫描型 PDF 暂不支持")
        return _assemble("pdf", pending)

    @staticmethod
    def _extract_docx(content: bytes) -> ExtractedDocument:
        if DocxDocument is None or CT_P is None or CT_Tbl is None:
            raise ExtractionError("当前运行环境缺少 DOCX 解析组件")
        try:
            document = DocxDocument(BytesIO(content))
            pending: list[_PendingBlock] = []
            heading_stack: list[str] = []
            paragraph_number = 0
            table_number = 0
            element_number = 0
            for child in document.element.body.iterchildren():
                if isinstance(child, CT_P):
                    paragraph_number += 1
                    paragraph = Paragraph(child, document)
                    text = "\n".join(_normalize_lines(paragraph.text)).strip("\n")
                    if not text.strip():
                        continue
                    element_number += 1
                    level = _heading_level(paragraph)
                    kind: BlockKind = "paragraph"
                    if level is not None:
                        heading_stack[:] = heading_stack[: level - 1]
                        heading_stack.append(text.strip())
                        kind = "heading"
                    pending.append(
                        _PendingBlock(
                            text=text,
                            kind=kind,
                            heading_path=tuple(heading_stack),
                            paragraph_start=paragraph_number,
                            paragraph_end=paragraph_number,
                            element_start=element_number,
                            element_end=element_number,
                        )
                    )
                elif isinstance(child, CT_Tbl):
                    table_number += 1
                    table = Table(child, document)
                    for row_number, row in enumerate(table.rows, start=1):
                        cells = [re.sub(r"\s*\n\s*", " / ", cell.text).strip() for cell in row.cells]
                        value = " | ".join(cells).strip(" |")
                        if not value:
                            continue
                        element_number += 1
                        pending.append(
                            _PendingBlock(
                                text=value,
                                kind="table_row",
                                heading_path=tuple(heading_stack),
                                table_start=table_number,
                                table_end=table_number,
                                table_row_start=row_number,
                                table_row_end=row_number,
                                element_start=element_number,
                                element_end=element_number,
                            )
                        )
        except Exception as exc:
            raise ExtractionError("DOCX 文件损坏或无法解析") from exc
        if not pending:
            raise ExtractionError("DOCX 没有可提取文本")
        return _assemble("docx", pending)
